"""Casos de uso de lectura de cuentas (E2-T02, HU05 CA-01/CA-02).

Lee la proyeccion `account_balances` via `accounts/repository` (E2-T01):
no calcula saldos por fuera del ledger (`docs/modules/README.md#accounts`);
`contable = disponible + retenido` (`03b#5.2`, invariante `03c#15.1`).

RBAC: el cliente solo ve sus cuentas; toda funcion recibe el `user_id` del
contexto de auth y filtra/verifica propiedad aqui (no en el router).
Enmascaramiento por defecto (regla de oro 9): el numero completo nunca sale
de este modulo; `mask_account_number` es dominio puro (sin BD).

Convencion: solo lectura (sin `flush`/`commit`); dinero entero en centimos,
nunca `float`.
"""

from __future__ import annotations

import csv
import io
import uuid
from datetime import date

from sqlalchemy.orm import Session

from app.modules.accounts import repository as accounts_repo
from app.modules.accounts.schemas import AccountDetail, AccountSummary, Movement

MASK_PREFIX = "****"

#: Paginacion de movimientos (E2-T03): defaults y tope (05#4 usa 20 de ejemplo).
MOVEMENTS_DEFAULT_PAGE = 1
MOVEMENTS_DEFAULT_PAGE_SIZE = 20
MOVEMENTS_MAX_PAGE_SIZE = 100
#: Techo de filas leidas de la proyeccion por consulta (E2-T04 `list_movements`
#: no expone filtros: se pagina/filtra en este caso de uso sobre esa lectura;
#: suficiente para el dataset de prueba razonable con <2 s; documentado).
MOVEMENTS_FETCH_LIMIT = 10_000
#: Techo de filas por export (rango de fechas obligatorio acota el volumen).
EXPORT_MAX_ROWS = 10_000

MOVEMENT_DIRECTIONS = ("DEBIT", "CREDIT")
EXPORT_FORMATS = ("csv", "xlsx", "pdf")
EXPORT_CSV_HEADERS = (
    "journal_entry_id",
    "transaction_id",
    "account_id",
    "direction",
    "amount_minor",
    "currency",
    "description",
    "value_date",
    "created_at",
)


class AccountNotFoundError(LookupError):
    """La cuenta no existe (el router lo traduce a 404 NOT_FOUND)."""


class AccountForbiddenError(PermissionError):
    """La cuenta no pertenece al usuario (el router lo traduce a 403 NOT_AUTHORIZED)."""


def mask_account_number(account_number: str) -> str:
    """Enmascara: solo ultimos 4 visibles, resto `*` (formato `****1234`).

    Dominio puro: sin BD, sin logs con PII (regla de oro 7: no se registra
    el numero completo en ningun lado).
    """
    if not isinstance(account_number, str) or not account_number:
        raise ValueError("account_number debe ser texto no vacio")
    last4 = account_number[-4:]
    return f"{MASK_PREFIX}{last4}"


def to_summary(account_id: uuid.UUID, **fields: object) -> AccountSummary:
    """Construye el consolidado de una cuenta (puro, sin BD).

    `balance_minor` (contable) = `available_minor + held_minor`, leidos de
    la proyeccion `account_balances` por el llamante.
    """
    available = int(fields["available_minor"])  # type: ignore[arg-type]
    held = int(fields["held_minor"])  # type: ignore[arg-type]
    if available < 0 or held < 0:
        raise ValueError("saldos de la proyeccion no pueden ser negativos")
    return AccountSummary(
        id=account_id,
        account_number_masked=mask_account_number(str(fields["account_number"])),
        type=str(fields["type"]),
        currency=str(fields["currency"]),
        available_minor=available,
        held_minor=held,
        balance_minor=available + held,
    )


def list_accounts(session: Session, user_id: uuid.UUID | str) -> list[AccountSummary]:
    """Consolidado del usuario: sus cuentas con saldos de la proyeccion."""
    try:
        uid = user_id if isinstance(user_id, uuid.UUID) else uuid.UUID(str(user_id))
    except (ValueError, AttributeError, TypeError) as exc:
        raise ValueError(f"user_id debe ser UUID, recibido: {user_id!r}") from exc
    items: list[AccountSummary] = []
    for account in accounts_repo.list_by_user(session, uid):
        balance = accounts_repo.get_balance(session, account.id)
        if balance is None:  # Sin proyeccion no hay saldo que exponer; se omite.
            continue
        items.append(
            to_summary(
                account.id,
                account_number=account.account_number,
                type=account.type,
                currency=balance.currency,
                available_minor=int(balance.available_minor),
                held_minor=int(balance.held_minor),
            )
        )
    return items


def get_account_detail(
    session: Session, user_id: uuid.UUID | str, account_id: uuid.UUID | str
) -> AccountDetail:
    """Detalle con disponible/retenido/contable; 404 si no existe, 403 si es ajena."""
    try:
        uid = user_id if isinstance(user_id, uuid.UUID) else uuid.UUID(str(user_id))
    except (ValueError, AttributeError, TypeError) as exc:
        raise ValueError(f"user_id debe ser UUID, recibido: {user_id!r}") from exc
    account = accounts_repo.get_account(session, account_id)
    if account is None:
        raise AccountNotFoundError(f"cuenta inexistente: {account_id}")
    if account.user_id != uid:
        raise AccountForbiddenError(f"cuenta {account.id} no pertenece al usuario")
    balance = accounts_repo.get_balance(session, account.id)
    if balance is None:
        raise AccountNotFoundError(f"cuenta sin proyeccion de saldos: {account.id}")
    summary = to_summary(
        account.id,
        account_number=account.account_number,
        type=account.type,
        currency=balance.currency,
        available_minor=int(balance.available_minor),
        held_minor=int(balance.held_minor),
    )
    return AccountDetail(**summary.model_dump(), status=account.status)


# ------------------------------------------------- Movimientos (E2-T03)


class ExportFormatUnavailableError(RuntimeError):
    """Formato de export no disponible (el router lo traduce a 501)."""


def openpyxl_available() -> bool:
    """`True` si `openpyxl` esta instalado (export xlsx real)."""
    import importlib.util

    return importlib.util.find_spec("openpyxl") is not None


def pdf_available() -> str | None:
    """Backend PDF disponible (`reportlab` > `fpdf` > `weasyprint`) o `None`.

    Prioridad documentada: `reportlab` (tabla platypus con `repeatRows`),
    luego `fpdf`/`fpdf2` (filas con salto de pagina manual), luego
    `weasyprint` (tabla HTML con `thead` repetido). Puro (solo sondea
    `importlib`, sin importar las libs).
    """
    import importlib.util

    for name in ("reportlab", "fpdf", "weasyprint"):
        if importlib.util.find_spec(name) is not None:
            return name
    return None


def to_movement(row: object) -> Movement:
    """Convierte una fila de `movements_view` al esquema de lectura (puro)."""
    get = getattr
    return Movement(
        journal_entry_id=get(row, "journal_entry_id"),
        transaction_id=get(row, "transaction_id"),
        account_id=get(row, "account_id"),
        direction=get(row, "direction"),
        amount_minor=int(get(row, "amount_minor")),  # type: ignore[arg-type]
        currency=get(row, "currency"),
        description=get(row, "description"),
        value_date=get(row, "value_date"),
        created_at=get(row, "created_at"),
    )


def _require_owned_account(session: Session, user_id: uuid.UUID | str, account_id: uuid.UUID | str):
    """Cuenta propia o error: 404 si no existe, 403 si es ajena (RBAC E2-T02)."""
    try:
        uid = user_id if isinstance(user_id, uuid.UUID) else uuid.UUID(str(user_id))
    except (ValueError, AttributeError, TypeError) as exc:
        raise ValueError(f"user_id debe ser UUID, recibido: {user_id!r}") from exc
    account = accounts_repo.get_account(session, account_id)
    if account is None:
        raise AccountNotFoundError(f"cuenta inexistente: {account_id}")
    if account.user_id != uid:
        raise AccountForbiddenError(f"cuenta {account.id} no pertenece al usuario")
    return account


def _validate_movement_filters(
    *,
    page: int,
    page_size: int,
    direction: str | None,
    date_from: date | None,
    date_to: date | None,
) -> None:
    if isinstance(page, bool) or not isinstance(page, int) or page < 1:
        raise ValueError(f"page debe ser int >= 1, recibido: {page!r}")
    if (
        isinstance(page_size, bool)
        or not isinstance(page_size, int)
        or page_size < 1
        or page_size > MOVEMENTS_MAX_PAGE_SIZE
    ):
        raise ValueError(
            f"page_size debe ser int 1..{MOVEMENTS_MAX_PAGE_SIZE}, recibido: {page_size!r}"
        )
    if direction is not None and direction not in MOVEMENT_DIRECTIONS:
        raise ValueError(f"direction debe ser DEBIT/CREDIT, recibido: {direction!r}")
    if date_from is not None and date_to is not None and date_from > date_to:
        raise ValueError(f"rango de fechas invertido: {date_from} > {date_to}")


def _apply_movement_filters(
    rows: list,
    *,
    date_from: date | None,
    date_to: date | None,
    direction: str | None,
) -> list:
    """Filtra filas ya leidas de `movements_view` (fecha valor + direccion).

    E2-T04 `list_movements` (reutilizado, sin duplicar SQL) no expone filtros:
    se filtra aqui en memoria sobre su lectura ordenada (recientes primero),
    preservando el orden para paginar despues. Solo lectura.
    """
    kept = []
    for row in rows:
        if direction is not None and row.direction != direction:
            continue
        value_date = row.value_date
        if date_from is not None and (value_date is None or value_date < date_from):
            continue
        if date_to is not None and (value_date is None or value_date > date_to):
            continue
        kept.append(row)
    return kept


def list_account_movements(
    session: Session,
    user_id: uuid.UUID | str,
    account_id: uuid.UUID | str,
    *,
    page: int = MOVEMENTS_DEFAULT_PAGE,
    page_size: int = MOVEMENTS_DEFAULT_PAGE_SIZE,
    date_from: date | None = None,
    date_to: date | None = None,
    direction: str | None = None,
) -> tuple[list[Movement], int]:
    """Movimientos paginados de una cuenta propia (HU05 CA-03).

    RBAC E2-T02: 404 si no existe, 403 si es ajena. Lee la proyeccion
    `movements_view` via E2-T04 `list_movements` (sin recalcular nada);
    filtra por fecha valor (`desde/hasta`) y `direction`; pagina con
    `page`/`page_size`. Retorna `(items_de_la_pagina, total_filtrado)`.
    Solo lectura (sin `flush`/`commit`); dinero entero en centimos.
    """
    _validate_movement_filters(
        page=page,
        page_size=page_size,
        direction=direction,
        date_from=date_from,
        date_to=date_to,
    )
    account = _require_owned_account(session, user_id, account_id)
    rows = accounts_repo.list_movements(session, account.id, limit=MOVEMENTS_FETCH_LIMIT, offset=0)
    filtered = _apply_movement_filters(
        rows, date_from=date_from, date_to=date_to, direction=direction
    )
    total = len(filtered)
    start = (page - 1) * page_size
    return [to_movement(r) for r in filtered[start : start + page_size]], total


def get_export_movements(
    session: Session,
    user_id: uuid.UUID | str,
    account_id: uuid.UUID | str,
    *,
    date_from: date,
    date_to: date,
    direction: str | None = None,
) -> list[Movement]:
    """Filas para export: exige rango de fechas (HU05 CA-04, MVP).

    `date_from`/`date_to` obligatorios (el router los declara requeridos;
    aqui se validan por uso directo tambien). Tope `EXPORT_MAX_ROWS`.
    """
    if date_from is None or date_to is None:
        raise ValueError("el export exige rango de fechas: date_from y date_to")
    _validate_movement_filters(
        page=1,
        page_size=MOVEMENTS_DEFAULT_PAGE_SIZE,
        direction=direction,
        date_from=date_from,
        date_to=date_to,
    )
    account = _require_owned_account(session, user_id, account_id)
    rows = accounts_repo.list_movements(session, account.id, limit=MOVEMENTS_FETCH_LIMIT, offset=0)
    filtered = _apply_movement_filters(
        rows, date_from=date_from, date_to=date_to, direction=direction
    )[:EXPORT_MAX_ROWS]
    return [to_movement(r) for r in filtered]


def _csv_safe(value: object) -> object:
    """Mitiga inyeccion de formulas CSV: prefija `' ` a `=,+,-,@` iniciales."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def movements_to_csv(rows: list[Movement]) -> bytes:
    """Serializa movimientos a CSV UTF-8 con BOM (abre bien en Excel).

    Cabeceras = columnas `03b#5.4` + `account_id`; montos como enteros
    (`amount_minor`), fechas ISO-8601. Puro (sin BD).
    """
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow(EXPORT_CSV_HEADERS)
    for m in rows:
        d = m.model_dump(mode="json")
        writer.writerow([_csv_safe(d.get(h)) for h in EXPORT_CSV_HEADERS])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def movements_to_xlsx(rows: list[Movement]) -> bytes:
    """Serializa movimientos a XLSX (`openpyxl`); 501 si no esta instalado.

    Decision E2-T03: `openpyxl` no esta en el `.venv` del backend, asi que el
    endpoint responde 501 `EXPORT_FORMAT_NOT_SUPPORTED` para `format=xlsx`
    hasta que la dependencia se agregue (CSV es el formato MVP garantizado).
    """
    if not openpyxl_available():
        raise ExportFormatUnavailableError(
            "export xlsx no disponible: openpyxl no instalado; use format=csv"
        )
    from openpyxl import Workbook  # import perezoso: dependencia opcional

    wb = Workbook()
    ws = wb.title and wb.active
    assert ws is not None
    ws.title = "movimientos"
    ws.append(list(EXPORT_CSV_HEADERS))
    for m in rows:
        d = m.model_dump(mode="json")
        ws.append([d.get(h) for h in EXPORT_CSV_HEADERS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def movements_to_pdf(
    rows: list[Movement],
    *,
    account_id: uuid.UUID | str,
    date_from: date,
    date_to: date,
) -> bytes:
    """Serializa movimientos a PDF (tabla `03b#5.4`); 501 si no hay backend.

    Decision fase 4/H1: ninguna lib PDF (`reportlab`/`fpdf2`/`weasyprint`)
    esta en el `.venv` del backend, asi que el endpoint responde 501
    `EXPORT_FORMAT_NOT_SUPPORTED` para `format=pdf` hasta que una se agregue
    (mismo patron que `movements_to_xlsx`). Cuando exista, genera PDF real:
    titulo con cuenta/rango, cabeceras `EXPORT_CSV_HEADERS`, montos como
    enteros (`amount_minor`, sin `float`), fechas ISO-8601 y paginado
    automatico (cabecera repetida por pagina). Solo UUIDs en movimientos
    (sin numeros en claro, regla de oro 9). Puro (sin BD).
    """
    backend = pdf_available()
    if backend is None:
        raise ExportFormatUnavailableError(
            "export pdf no disponible: instale reportlab, fpdf2 o weasyprint; " "use format=csv"
        )
    table_rows = [[str(m.model_dump(mode="json").get(h)) for h in EXPORT_CSV_HEADERS] for m in rows]
    title = f"Movimientos {account_id} {date_from.isoformat()} a {date_to.isoformat()}"
    if backend == "reportlab":
        return _movements_to_pdf_reportlab(table_rows, title)
    if backend == "fpdf":
        return _movements_to_pdf_fpdf(table_rows, title)
    return _movements_to_pdf_weasyprint(table_rows, title)


def _movements_to_pdf_reportlab(table_rows: list[list[str]], title: str) -> bytes:
    """PDF real con `reportlab` platypus (paginado via `repeatRows=1`)."""
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4))
    styles = getSampleStyleSheet()
    story = [Paragraph(escape(title), styles["Heading2"]), Spacer(1, 12)]
    data = [list(EXPORT_CSV_HEADERS)] + table_rows
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return buf.getvalue()


def _movements_to_pdf_fpdf(table_rows: list[list[str]], title: str) -> bytes:
    """PDF real con `fpdf2` (cabecera repetida por pagina, salto manual)."""
    from fpdf import FPDF

    pdf = FPDF(orientation="L", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    col_w = 31  # 9 cols x 31mm aprox. en A4 apaisado con margenes
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "B", 7)
    for h in EXPORT_CSV_HEADERS:
        pdf.cell(col_w, 8, h, border=1)
    pdf.ln()
    pdf.set_font("Helvetica", "", 7)
    for row in table_rows:
        if pdf.get_y() > 170:
            pdf.add_page()
            pdf.set_font("Helvetica", "B", 7)
            for h in EXPORT_CSV_HEADERS:
                pdf.cell(col_w, 8, h, border=1)
            pdf.ln()
            pdf.set_font("Helvetica", "", 7)
        for cell in row:
            pdf.cell(col_w, 6, (cell or "")[:40], border=1)
        pdf.ln()
    raw = pdf.output()
    return bytes(raw) if isinstance(raw, (bytes, bytearray)) else raw.encode("latin-1")


def _movements_to_pdf_weasyprint(table_rows: list[list[str]], title: str) -> bytes:
    """PDF real con `weasyprint` (tabla HTML, `thead` repetido por pagina)."""
    from html import escape

    from weasyprint import HTML

    cells = "".join(f"<th>{escape(h)}</th>" for h in EXPORT_CSV_HEADERS)
    body = "".join(
        "<tr>" + "".join(f"<td>{escape(c or '')}</td>" for c in row) + "</tr>" for row in table_rows
    )
    html = (
        "<html><head><meta charset='utf-8'>"
        "<style>@page{size:A4 landscape;@bottom-center{content:counter(page)}}"
        "table{border-collapse:collapse;width:100%;font-size:8px}"
        "th,td{border:1px solid #888;padding:2px}thead{display:table-header-group}"
        "</style></head><body>"
        f"<h2>{escape(title)}</h2>"
        f"<table><thead><tr>{cells}</tr></thead><tbody>{body}</tbody></table>"
        "</body></html>"
    )
    return HTML(string=html).write_pdf()


__all__ = [
    "EXPORT_CSV_HEADERS",
    "EXPORT_FORMATS",
    "EXPORT_MAX_ROWS",
    "MASK_PREFIX",
    "MOVEMENTS_DEFAULT_PAGE",
    "MOVEMENTS_DEFAULT_PAGE_SIZE",
    "MOVEMENTS_FETCH_LIMIT",
    "MOVEMENTS_MAX_PAGE_SIZE",
    "MOVEMENT_DIRECTIONS",
    "AccountForbiddenError",
    "AccountNotFoundError",
    "ExportFormatUnavailableError",
    "get_account_detail",
    "get_export_movements",
    "list_account_movements",
    "list_accounts",
    "mask_account_number",
    "movements_to_csv",
    "movements_to_pdf",
    "movements_to_xlsx",
    "openpyxl_available",
    "pdf_available",
    "to_movement",
    "to_summary",
]
